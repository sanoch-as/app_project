from datetime import date, timedelta

import pytest
from httpx import AsyncClient

from tests.fixtures.auth import auth_header, invite_member, register_admin

pytestmark = pytest.mark.asyncio


async def _create_project(
    client: AsyncClient, admin_token: str, name: str = "Dashboard Test"
) -> str:
    response = await client.post(
        "/api/v1/projects", json={"name": name}, headers=auth_header(admin_token)
    )
    return response.json()["id"]


async def _create_task(client: AsyncClient, admin_token: str, project_id: str, **overrides) -> dict:
    payload = {
        "name": "Task",
        "start_date": "2026-09-01",
        "duration_days": 5,
        "budgeted_cost": 1000,
    }
    payload.update(overrides)
    response = await client.post(
        f"/api/v1/projects/{project_id}/tasks", json=payload, headers=auth_header(admin_token)
    )
    assert response.status_code == 201, response.text
    return response.json()


async def test_project_dashboard_reports_percent_complete_and_overdue(client: AsyncClient):
    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    project_id = await _create_project(client, admin_token)

    overdue_date = (date.today() - timedelta(days=10)).isoformat()
    task = await _create_task(
        client, admin_token, project_id, start_date=overdue_date, duration_days=1
    )
    await client.patch(
        f"/api/v1/tasks/{task['id']}",
        json={"percent_complete": 50},
        headers=auth_header(admin_token),
    )

    response = await client.get(
        f"/api/v1/projects/{project_id}/dashboard", headers=auth_header(admin_token)
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["percent_complete"] == 50.0
    assert len(body["overdue_tasks"]) == 1
    assert body["overdue_tasks"][0]["id"] == task["id"]


async def test_dashboard_percent_complete_excludes_wbs_parent_rollup(client: AsyncClient):
    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    project_id = await _create_project(client, admin_token)

    parent = await _create_task(client, admin_token, project_id, name="Parent")
    child_one = await _create_task(
        client,
        admin_token,
        project_id,
        name="Child 1",
        parent_task_id=parent["id"],
        budgeted_cost=800,
    )
    await client.patch(
        f"/api/v1/tasks/{child_one['id']}",
        json={"percent_complete": 100},
        headers=auth_header(admin_token),
    )
    await _create_task(
        client,
        admin_token,
        project_id,
        name="Child 2",
        parent_task_id=parent["id"],
        budgeted_cost=200,
    )
    # Unrelated top-level task, untouched (0% complete) — needed so a bug
    # that double-counts the parent's rolled-up duration/percent alongside
    # its children would actually shift the ratio (with only the
    # parent/children group in the project, doubling both sides cancels out).
    await _create_task(client, admin_token, project_id, name="Unrelated")

    response = await client.get(
        f"/api/v1/projects/{project_id}/dashboard", headers=auth_header(admin_token)
    )
    assert response.status_code == 200, response.text
    # Leaf-only, duration-weighted average (all three leaves have the same
    # 5-day duration, so this is just a plain average): (100+0+0)/3 = 33.33%.
    # Counting the parent's own rolled-up duration/percent on top would
    # double its children's contribution and skew this to 37.5%.
    assert response.json()["percent_complete"] == pytest.approx(100 / 3)


async def test_project_dashboard_lists_upcoming_milestones(client: AsyncClient):
    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    project_id = await _create_project(client, admin_token)

    future_date = (date.today() + timedelta(days=30)).isoformat()
    milestone = await _create_task(
        client,
        admin_token,
        project_id,
        name="Launch",
        start_date=future_date,
        duration_days=0,
        is_milestone=True,
    )

    response = await client.get(
        f"/api/v1/projects/{project_id}/dashboard", headers=auth_header(admin_token)
    )
    body = response.json()
    assert len(body["upcoming_milestones"]) == 1
    assert body["upcoming_milestones"][0]["id"] == milestone["id"]


async def test_portfolio_summary_lists_visible_projects_only(client: AsyncClient):
    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    member = await invite_member(client, admin_token=admin_token)

    visible_project = await _create_project(client, admin_token, "Visible")
    await _create_project(client, admin_token, "Hidden")
    await client.post(
        f"/api/v1/projects/{visible_project}/members",
        json={"user_id": member["user"]["id"]},
        headers=auth_header(admin_token),
    )

    admin_summary = await client.get("/api/v1/dashboard/summary", headers=auth_header(admin_token))
    assert admin_summary.json()["total_projects"] == 2

    member_summary = await client.get(
        "/api/v1/dashboard/summary", headers=auth_header(member["tokens"]["access_token"])
    )
    assert member_summary.json()["total_projects"] == 1
    assert member_summary.json()["projects"][0]["id"] == visible_project


async def test_export_tasks_csv(client: AsyncClient):
    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    project_id = await _create_project(client, admin_token)
    await _create_task(client, admin_token, project_id, name="Exportable Task")

    response = await client.get(
        f"/api/v1/projects/{project_id}/reports/export?type=tasks",
        headers=auth_header(admin_token),
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "Exportable Task" in response.text
    assert "wbs_code" in response.text.splitlines()[0]


async def test_export_tasks_xlsx_is_styled_and_hierarchical(client: AsyncClient):
    import io

    from openpyxl import load_workbook

    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    project_id = await _create_project(client, admin_token)

    parent = await _create_task(client, admin_token, project_id, name="Fase 1")
    await _create_task(
        client,
        admin_token,
        project_id,
        name="Subtarea 1.1",
        parent_task_id=parent["id"],
        priority="critical",
    )

    response = await client.get(
        f"/api/v1/projects/{project_id}/reports/export?type=tasks_xlsx",
        headers=auth_header(admin_token),
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].endswith('.xlsx"')

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    flattened = [v for row in values for v in row if isinstance(v, str)]
    assert "Fase 1" in flattened
    assert "Subtarea 1.1" in flattened
    assert "Clave" in values[3]  # header row (row 4: title, subtitle, blank, header)

    # The child's own row has its priority cell ("Crítica") filled with the
    # same red used for the "critical" badge in the web app.
    header_row_idx = next(i for i, row in enumerate(values, start=1) if "Clave" in row)
    child_row_idx = next(i for i, row in enumerate(values, start=1) if "Subtarea 1.1" in row)
    priority_col = values[header_row_idx - 1].index("Prioridad") + 1
    priority_cell = sheet.cell(row=child_row_idx, column=priority_col)
    assert priority_cell.value == "Crítica"
    assert priority_cell.fill.fgColor.rgb == "00FEF2F2"


async def test_export_tasks_pdf_is_a_real_pdf(client: AsyncClient):
    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    project_id = await _create_project(client, admin_token)
    await _create_task(client, admin_token, project_id, name="Exportable Task")

    response = await client.get(
        f"/api/v1/projects/{project_id}/reports/export?type=tasks_pdf",
        headers=auth_header(admin_token),
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].endswith('.pdf"')
    assert response.content.startswith(b"%PDF")


async def test_export_worklogs_csv(client: AsyncClient):
    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    project_id = await _create_project(client, admin_token)
    task = await _create_task(client, admin_token, project_id)
    await client.post(
        f"/api/v1/tasks/{task['id']}/worklogs",
        json={"work_date": "2026-09-02", "hours": 3, "description": "Worked on it"},
        headers=auth_header(admin_token),
    )

    response = await client.get(
        f"/api/v1/projects/{project_id}/reports/export?type=worklogs",
        headers=auth_header(admin_token),
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "Worked on it" in response.text


async def test_export_summary_pdf(client: AsyncClient):
    admin = await register_admin(client)
    admin_token = admin["tokens"]["access_token"]
    project_id = await _create_project(client, admin_token)
    await _create_task(client, admin_token, project_id)
    await client.post(
        f"/api/v1/projects/{project_id}/baselines",
        json={"name": "Plan"},
        headers=auth_header(admin_token),
    )

    response = await client.get(
        f"/api/v1/projects/{project_id}/reports/export?type=summary",
        headers=auth_header(admin_token),
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
