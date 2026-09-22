"""CSV/PDF export — spec section 4.1 point 24. Pure formatting functions:
callers (endpoints/reports.py) load the data and hand it to these."""

import csv
import io
import uuid
import xml.sax.saxutils
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.enums import DateFormat, Language, TaskPriority, TaskStatus
from app.models.project import Project
from app.models.task import Task
from app.repositories.worklog_repository import WorklogExportRow
from app.services.evm import EVMMetrics
from app.services.scurve import SCurvePoint

TASKS_CSV_HEADER = [
    "wbs_code",
    "name",
    "status",
    "priority",
    "start_date",
    "end_date",
    "duration_days",
    "percent_complete",
    "is_milestone",
    "is_critical",
    "total_float",
    "budgeted_cost",
]

WORKLOGS_CSV_HEADER = ["work_date", "user_email", "task_name", "hours", "description"]


def tasks_to_csv(tasks: list[Task]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(TASKS_CSV_HEADER)
    for task in tasks:
        writer.writerow(
            [
                task.wbs_code,
                task.name,
                task.status.value,
                task.priority.value,
                task.start_date.isoformat(),
                task.end_date.isoformat(),
                task.duration_days,
                task.percent_complete,
                task.is_milestone,
                task.is_critical,
                task.total_float if task.total_float is not None else "",
                task.budgeted_cost,
            ]
        )
    return buffer.getvalue()


def worklogs_to_csv(rows: list[WorklogExportRow]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(WORKLOGS_CSV_HEADER)
    for row in rows:
        writer.writerow(
            [
                row.work_date.isoformat(),
                row.user_email,
                row.task_name,
                row.hours,
                row.description or "",
            ]
        )
    return buffer.getvalue()


def project_summary_to_pdf(
    project: Project, metrics: EVMMetrics, scurve: list[SCurvePoint], as_of: date
) -> bytes:
    """A simple one-page PDF: project header, KPI table, and the S-curve's
    numbers as a table (section 4.1 point 24: "PDF simple del resumen de
    proyecto — curva S + KPIs"; no chart rendering, just the figures)."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"Project Summary — {project.name}", styles["Title"]),
        Paragraph(f"Status: {project.status.value} — As of: {as_of.isoformat()}", styles["Normal"]),
        Spacer(1, 0.25 * inch),
        Paragraph("Key Performance Indicators", styles["Heading2"]),
    ]

    kpi_data = [
        ["Metric", "Value"],
        ["Planned Value (PV)", f"{metrics.pv:,.2f}"],
        ["Earned Value (EV)", f"{metrics.ev:,.2f}"],
        ["Actual Cost (AC)", f"{metrics.ac:,.2f}"],
        ["SPI", f"{metrics.spi:.2f}" if metrics.spi is not None else "N/A"],
        ["CPI", f"{metrics.cpi:.2f}" if metrics.cpi is not None else "N/A"],
    ]
    story.append(_styled_table(kpi_data))
    story.append(Spacer(1, 0.25 * inch))

    story.append(Paragraph("S-Curve (weekly)", styles["Heading2"]))
    if scurve:
        curve_data = [["Week Ending", "PV", "EV", "AC", "SPI", "CPI"]]
        for point in scurve:
            curve_data.append(
                [
                    point.week_ending.isoformat(),
                    f"{point.metrics.pv:,.2f}",
                    f"{point.metrics.ev:,.2f}",
                    f"{point.metrics.ac:,.2f}",
                    f"{point.metrics.spi:.2f}" if point.metrics.spi is not None else "N/A",
                    f"{point.metrics.cpi:.2f}" if point.metrics.cpi is not None else "N/A",
                ]
            )
        story.append(_styled_table(curve_data))
    else:
        story.append(Paragraph("No baseline saved yet — nothing to plot.", styles["Normal"]))

    doc.build(story)
    return buffer.getvalue()


def _styled_table(data: list[list[str]]) -> Table:
    table = Table(data)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E3440")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F0F0")]),
            ]
        )
    )
    return table


# ---------------------------------------------------------------------------
# Styled task-list export (Tasks tab "Descargar" button) — Excel and PDF,
# matching the Tasks tab's own tree view: same row order (WBS hierarchy,
# numerically sorted, not the raw string sort docs/BACKLOG.md notes as a
# known limitation elsewhere), same indentation, same status/priority badge
# colors (mirrors frontend/src/components/common/Badge.tsx + tailwind
# config's `jira` palette, kept in sync by hand — the two apps don't share a
# module). Built to be handed straight to a client: title, generation date,
# frozen header, no gridlines, landscape/fit-to-width page setup.
# ---------------------------------------------------------------------------

_TASK_EXPORT_HEADERS: dict[str, list[str]] = {
    "es": ["Clave", "Nombre", "Estado", "Prioridad", "Inicio", "Fin", "% Hecho"],
    "en": ["Key", "Name", "Status", "Priority", "Start", "End", "% Done"],
}

_STATUS_LABELS: dict[str, dict[TaskStatus, str]] = {
    "es": {
        TaskStatus.NOT_STARTED: "Sin iniciar",
        TaskStatus.IN_PROGRESS: "En curso",
        TaskStatus.BLOCKED: "Bloqueada",
        TaskStatus.COMPLETED: "Completada",
    },
    "en": {
        TaskStatus.NOT_STARTED: "Not started",
        TaskStatus.IN_PROGRESS: "In progress",
        TaskStatus.BLOCKED: "Blocked",
        TaskStatus.COMPLETED: "Completed",
    },
}

_PRIORITY_LABELS: dict[str, dict[TaskPriority, str]] = {
    "es": {
        TaskPriority.LOW: "Baja",
        TaskPriority.MEDIUM: "Media",
        TaskPriority.HIGH: "Alta",
        TaskPriority.CRITICAL: "Crítica",
    },
    "en": {
        TaskPriority.LOW: "Low",
        TaskPriority.MEDIUM: "Medium",
        TaskPriority.HIGH: "High",
        TaskPriority.CRITICAL: "Critical",
    },
}

# (background, text) hex pairs, no "#" — matches Badge.tsx's tone classes.
_STATUS_COLORS: dict[TaskStatus, tuple[str, str]] = {
    TaskStatus.NOT_STARTED: ("F1F2F4", "44546F"),  # gray
    TaskStatus.IN_PROGRESS: ("E9F2FF", "0C66E4"),  # blue
    TaskStatus.BLOCKED: ("FEF2F2", "E2483D"),  # red
    TaskStatus.COMPLETED: ("DCFFF1", "216E4E"),  # green
}
_PRIORITY_COLORS: dict[TaskPriority, tuple[str, str]] = {
    TaskPriority.LOW: ("F1F2F4", "44546F"),  # gray
    TaskPriority.MEDIUM: ("E9F2FF", "0C66E4"),  # blue
    TaskPriority.HIGH: ("FFF7ED", "E56910"),  # orange
    TaskPriority.CRITICAL: ("FEF2F2", "E2483D"),  # red
}

_HEADER_BG = "FAFBFC"  # jira.panel
_TEXT = "172B4D"  # jira.text
_TEXT_SUB = "626F86"  # jira.textSub
_BORDER = "DCDFE4"  # jira.border
_STRIPE_BG = "FAFBFC"  # jira.panel, used as an alternating-row tint


def _export_language(language: Language) -> str:
    return language.value if language in (Language.ES, Language.EN) else "es"


def _format_export_date(value: date, date_format: DateFormat) -> str:
    return value.strftime("%d/%m/%Y") if date_format == DateFormat.DMY else value.isoformat()


def _wbs_sort_key(wbs_code: str) -> tuple[int, ...]:
    """Numeric, not lexicographic — "10" must sort after "2", not before it
    (docs/BACKLOG.md's known WBS-ordering limitation elsewhere in the app)."""
    parts: list[int] = []
    for segment in wbs_code.split("."):
        try:
            parts.append(int(segment))
        except ValueError:
            parts.append(0)
    return tuple(parts)


def _build_hierarchical_rows(tasks: list[Task]) -> list[tuple[Task, int]]:
    """Depth-first WBS order (a parent immediately followed by its own
    subtree, siblings numerically sorted) — the same shape as the Tasks
    tab's own tree view, so the exported row order matches what's on
    screen."""
    children_by_parent: dict[uuid.UUID | None, list[Task]] = {}
    for task in tasks:
        children_by_parent.setdefault(task.parent_task_id, []).append(task)
    for siblings in children_by_parent.values():
        siblings.sort(key=lambda t: _wbs_sort_key(t.wbs_code))

    rows: list[tuple[Task, int]] = []

    def visit(parent_id: uuid.UUID | None, depth: int) -> None:
        for task in children_by_parent.get(parent_id, []):
            rows.append((task, depth))
            visit(task.id, depth + 1)

    visit(None, 0)
    return rows


def tasks_to_styled_xlsx(
    project: Project, tasks: list[Task], language: Language, date_format: DateFormat
) -> bytes:
    lang = _export_language(language)
    headers = _TASK_EXPORT_HEADERS[lang]
    status_labels = _STATUS_LABELS[lang]
    priority_labels = _PRIORITY_LABELS[lang]
    rows = _build_hierarchical_rows(tasks)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (project.name or "Tareas")[:31] or "Tareas"

    thin_border = Border(*([Side(style="thin", color=_BORDER)] * 4))

    title_text = f"{project.name} — {'Listado de tareas' if lang == 'es' else 'Task list'}"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(bold=True, size=14, color=_TEXT)
    sheet.row_dimensions[1].height = 26

    generated_prefix = "Generado el " if lang == "es" else "Generated on "
    subtitle_text = generated_prefix + _format_export_date(date.today(), date_format)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    subtitle_cell = sheet.cell(row=2, column=1, value=subtitle_text)
    subtitle_cell.font = Font(size=9, italic=True, color=_TEXT_SUB)

    header_row = 4
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, size=10, color=_TEXT)
        cell.fill = PatternFill("solid", fgColor=_HEADER_BG)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for offset, (task, depth) in enumerate(rows):
        row_idx = header_row + 1 + offset
        striped = offset % 2 == 1
        values = [
            task.wbs_code,
            task.name,
            status_labels.get(task.status, task.status.value),
            priority_labels.get(task.priority, task.priority.value),
            _format_export_date(task.start_date, date_format),
            _format_export_date(task.end_date, date_format),
            f"{float(task.percent_complete):.1f}%",
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.font = Font(size=10, color=_TEXT)
            if striped:
                cell.fill = PatternFill("solid", fgColor=_STRIPE_BG)
            cell.alignment = (
                Alignment(indent=depth * 2, vertical="center")
                if col == 2
                else Alignment(vertical="center")
            )

        status_bg, status_fg = _STATUS_COLORS.get(task.status, ("FFFFFF", _TEXT))
        status_cell = sheet.cell(row=row_idx, column=3)
        status_cell.fill = PatternFill("solid", fgColor=status_bg)
        status_cell.font = Font(size=10, bold=True, color=status_fg)

        priority_bg, priority_fg = _PRIORITY_COLORS.get(task.priority, ("FFFFFF", _TEXT))
        priority_cell = sheet.cell(row=row_idx, column=4)
        priority_cell.fill = PatternFill("solid", fgColor=priority_bg)
        priority_cell.font = Font(size=10, bold=True, color=priority_fg)

    for col, width in enumerate([12, 42, 16, 14, 13, 13, 11], start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def tasks_to_styled_pdf(
    project: Project, tasks: list[Task], language: Language, date_format: DateFormat
) -> bytes:
    lang = _export_language(language)
    headers = _TASK_EXPORT_HEADERS[lang]
    status_labels = _STATUS_LABELS[lang]
    priority_labels = _PRIORITY_LABELS[lang]
    rows = _build_hierarchical_rows(tasks)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.4 * inch,
        rightMargin=0.4 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.4 * inch,
    )
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("TaskExportCell", parent=styles["Normal"], fontSize=8, leading=10)

    title = "Listado de tareas" if lang == "es" else "Task list"
    generated_prefix = "Generado el " if lang == "es" else "Generated on "
    generated_text = generated_prefix + _format_export_date(date.today(), date_format)
    story = [
        Paragraph(f"{xml.sax.saxutils.escape(project.name)} — {title}", styles["Title"]),
        Paragraph(generated_text, styles["Normal"]),
        Spacer(1, 0.2 * inch),
    ]

    data: list[list[object]] = [list(headers)]
    for task, depth in rows:
        indent = "&nbsp;" * (depth * 4)
        name_cell = Paragraph(f"{indent}{xml.sax.saxutils.escape(task.name)}", cell_style)
        data.append(
            [
                task.wbs_code,
                name_cell,
                status_labels.get(task.status, task.status.value),
                priority_labels.get(task.priority, task.priority.value),
                _format_export_date(task.start_date, date_format),
                _format_export_date(task.end_date, date_format),
                f"{float(task.percent_complete):.1f}%",
            ]
        )

    col_widths = [
        0.7 * inch,
        3.6 * inch,
        1.1 * inch,
        0.9 * inch,
        0.9 * inch,
        0.9 * inch,
        0.8 * inch,
    ]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    style_commands: list[tuple[object, ...]] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEADER_BG}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{_TEXT}")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{_BORDER}")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{_STRIPE_BG}")]),
    ]
    for row_idx, (task, _depth) in enumerate(rows, start=1):
        status_bg, status_fg = _STATUS_COLORS.get(task.status, ("FFFFFF", _TEXT))
        style_commands.append(
            ("BACKGROUND", (2, row_idx), (2, row_idx), colors.HexColor(f"#{status_bg}"))
        )
        style_commands.append(
            ("TEXTCOLOR", (2, row_idx), (2, row_idx), colors.HexColor(f"#{status_fg}"))
        )
        priority_bg, priority_fg = _PRIORITY_COLORS.get(task.priority, ("FFFFFF", _TEXT))
        style_commands.append(
            ("BACKGROUND", (3, row_idx), (3, row_idx), colors.HexColor(f"#{priority_bg}"))
        )
        style_commands.append(
            ("TEXTCOLOR", (3, row_idx), (3, row_idx), colors.HexColor(f"#{priority_fg}"))
        )
    table.setStyle(TableStyle(style_commands))
    story.append(table)

    doc.build(story)
    return buffer.getvalue()
